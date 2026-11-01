import sys, datetime as dt
from openpyxl import Workbook
out=sys.argv[1]
# 1 invoice: title, blank, header at row 3, formulas, total row
wb=Workbook(); ws=wb.active; ws.title="Invoice"
ws["A1"]="Invoice 2026-10 — Acme"; ws.merge_cells("A1:D1")
ws.append([]); ws.append(["Item","Qty","Unit price","Amount"])
items=[("Paper",10,3.5),("Toner",2,48),("Pens",30,0.8),("Stapler",1,12.9),("Folders",25,1.2)]
for i,(n,q,p) in enumerate(items, start=4):
    ws.append([n,q,p,f"=B{i}*C{i}"]); ws[f"C{i}"].number_format='"$"#,##0.00'; ws[f"D{i}"].number_format='"$"#,##0.00'
ws.append(["Total",None,None,"=SUM(D4:D8)"]); ws["D9"].number_format='"$"#,##0.00'
wb.save(f"{out}/invoice.xlsx")
# 2 household: Jan, Feb, Summary with cross-sheet refs
wb=Workbook(); first=True
for m in ["Jan","Feb"]:
    ws=wb.active if first else wb.create_sheet(); first=False; ws.title=m
    ws.append(["Date","Category","Amount","Paid"])
    for d in range(1,9):
        ws.append([dt.date(2026,1 if m=="Jan" else 2,d*3),["Food","Rent","Transport"][d%3],round(12.5*d,2),d%2==0])
        ws.cell(row=ws.max_row,column=1).number_format="yyyy/m/d"
    ws.append([None,"Total","=SUM(C2:C9)"])
ws=wb.create_sheet("Summary"); ws.append(["Month","Spent"]); ws.append(["Jan","=Jan!C10"]); ws.append(["Feb","=Feb!C10"]); ws.append(["Average","=AVERAGE(B2:B3)"])
wb.save(f"{out}/household.xlsx")
# 3 contacts: leading zeros text, percentages, dates
wb=Workbook(); ws=wb.active; ws.title="Contacts"
ws.append(["Name","Phone","Joined","Discount","Score"])
for i,n in enumerate(["王小明","佐藤花子","Alice","Bob"]):
    ws.append([n,f"09{i}1234567",dt.date(2025,i+1,15),0.05*(i+1),80+i])
    ws.cell(row=ws.max_row,column=2).number_format="@"; ws.cell(row=ws.max_row,column=4).number_format="0%"
ws.append(["=\"Count: \"&COUNTA(A2:A5)",None,None,"=AVERAGE(D2:D5)","=IF(E2>81,\"high\",\"low\")"])
wb.save(f"{out}/contacts.xlsx")
